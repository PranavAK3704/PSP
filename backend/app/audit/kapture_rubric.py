"""Kapture-audit rubric — a DEDICATED, versioned quality rubric for auditing external
Kapture CRM tickets, kept separate from the internal-resolution rubric (audit/rubric.py)
so the two never interfere.

A QA lead authors HOW a Kapture resolution should be judged: which factors matter and how
heavily. The rubric is data, not code — editing it (or uploading a QA-guidelines document
that the machine structures) re-versions it, and the Kapture judge (audit/kapture.py) scores
each ticket against the current version, stamping every audit with its rubric_version.

Mirrors audit/rubric.py exactly (same store/versioning idiom) + a structure-from-text upload
path cloned from knowledge/governance.py::structure_framework_from_text.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import threading
from datetime import datetime, timezone

from ..durable_state import durable_path
from ..llm import registry as llm_registry
from ..llm.gemini_provider import _parse_json
from .rubric import _clean_dimensions   # reuse the internal rubric's dimension normaliser

# MUTABLE authored content → durable state dir (survives redeploys); default backend/data.
_STORE = durable_path("kapture_rubric.json")
_lock = threading.Lock()


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


# ── Seed rubric (version 1) — the REAL Valmo BAU audit model (from "BAU Audits_Aspire",
# the sheet Manju & Aarti audit against). Three tiers, encoded by KEY PREFIX (tier_of):
#   • QUALITY  (no prefix)  — scored Pass/Fail; `weight` is the POINT value; the 7 quality
#                             weights sum to 100 (Audit Score = % of scorable points passed).
#   • zt_…     zero-tolerance — an AUTO-FAIL gate (weight 0, carries no points).
#   • fatal_…  fatal          — an AUTO-FAIL gate (weight 0, carries no points).
# Any zt_/fatal_ violation collapses the audit to composite 0 / status FAIL, regardless of the
# quality points. The judge also returns a 0–1 gradient per factor (kept for coaching only).
_SEED_DIMENSIONS = [
    # ── QUALITY (Pass/Fail, points sum to 100) ──
    {"key": "proper_opening_closing", "label": "Proper opening & closing", "weight": 4.0,
     "description": "Proper salutation in the opening (e.g. \"Dear Partner,\") and a complete closing "
                    "(e.g. \"Thank you, Valmo Partner Support\")."},
    {"key": "correct_email_format", "label": "Correct email format", "weight": 6.0,
     "description": "Correct font / size / alignment and paragraph formatting (Arial/Times/Calibri 10, "
                    "left-aligned); no spacing, case, spelling or punctuation errors."},
    {"key": "empathy_acknowledgement", "label": "Empathy / apology / reassurance / acknowledgement", "weight": 10.0,
     "description": "Appropriate empathy, apology, reassurance and acknowledgement; the partner's concern is "
                    "paraphrased/acknowledged; rebuttal used where warranted."},
    {"key": "simple_language", "label": "Simple, easy to understand", "weight": 12.5,
     "description": "Grammatically correct, simple language; no company jargon, complex words or unclear statements."},
    {"key": "template_modification", "label": "Appropriate template modification", "weight": 25.0,
     "description": "The template is not modified unnecessarily or inappropriately (non-critical unless the "
                    "resolution itself is impacted)."},
    {"key": "email_flow", "label": "Adhered to email flow", "weight": 12.5,
     "description": "The prescribed email flow was followed in the correct order."},
    {"key": "app_education", "label": "Educated to use the app", "weight": 30.0,
     "description": "Offered the relevant self-help option / existing app feature proactively. Mark NA when the "
                    "issue has no applicable app self-help."},
    # ── ZERO-TOLERANCE gates (auto-fail; no points) ──
    {"key": "zt_language", "label": "ZT · Rude / sarcastic / abusive language", "weight": 0.0,
     "description": "ZERO-TOLERANCE gate. Fails ONLY if the agent used casual, sarcastic, rude or abusive "
                    "language. Any breach auto-fails the whole audit."},
    {"key": "zt_financial_loss", "label": "ZT · Action leading to financial loss", "weight": 0.0,
     "description": "ZERO-TOLERANCE gate. Fails ONLY if the agent took an action (e.g. wrong validation) that "
                    "causes a financial loss. Any breach auto-fails the whole audit."},
    # ── FATAL gates (auto-fail; no points) ──
    {"key": "fatal_crm_utilization", "label": "Fatal · Improper CRM utilization", "weight": 0.0,
     "description": "FATAL gate. Fails ONLY on clear CRM misuse: past tickets not checked in Kapture, wrong "
                    "notes/email ID, duplicate tickets not merged, or mandatory details not captured."},
    {"key": "fatal_incorrect_reversal", "label": "Fatal · Incorrect reversal request", "weight": 0.0,
     "description": "FATAL gate. Fails ONLY if the agent raised a wrong reversal request, or failed to raise "
                    "one that was clearly required."},
    {"key": "fatal_tagging", "label": "Fatal · Incorrect / no tagging", "weight": 0.0,
     "description": "FATAL gate. Fails ONLY on clearly incorrect (or missing) disposition/folder tagging or "
                    "ticket status."},
    {"key": "fatal_assignment", "label": "Fatal · Incorrect ticket assignment / handling", "weight": 0.0,
     "description": "FATAL gate. Fails ONLY if the ticket was assigned to the wrong queue or mishandled."},
    {"key": "fatal_misleading_info", "label": "Fatal · Incorrect / misleading information", "weight": 0.0,
     "description": "FATAL gate. Fails ONLY if the agent gave incorrect or misleading information — wrong TAT "
                    "vs SOP, wrong resolution/expectations, or wrong template customization."},
    {"key": "fatal_incomplete_info", "label": "Fatal · Incomplete information", "weight": 0.0,
     "description": "FATAL gate. Fails ONLY if the agent left a raised query unaddressed or gave a clearly "
                    "incomplete resolution."},
]


_SEED_FP = hashlib.md5(json.dumps(_SEED_DIMENSIONS, sort_keys=True).encode()).hexdigest()[:12]


def tier_of(key: str) -> str:
    """A dimension's criticality tier, encoded by key prefix. zt_/fatal_ are auto-fail gates;
    everything else is a scored quality parameter (its `weight` is its point value)."""
    k = (key or "").lower()
    if k.startswith("zt_"):
        return "zt"
    if k.startswith("fatal_"):
        return "fatal"
    return "quality"


def _seed() -> dict:
    return {"version": 1, "seed_fp": _SEED_FP,
            "dimensions": [dict(d) for d in _SEED_DIMENSIONS], "updated_at": _now()}


def _load() -> dict | None:
    if _STORE.exists():
        try:
            data = json.loads(_STORE.read_text())
            if isinstance(data, dict) and data.get("dimensions"):
                return data
        except Exception:  # noqa: BLE001
            return None
    return None


def _write(rubric: dict) -> None:
    _STORE.write_text(json.dumps(rubric, indent=1))


def get_rubric() -> dict:
    """Current Kapture rubric (seeds the QA-factor version 1 on first run)."""
    with _lock:
        rubric = _load()
        if rubric is None:
            rubric = _seed()
            _write(rubric)
            return rubric
        # Auto-refresh a NEVER-EDITED (version 1) rubric whenever the seed changes — by
        # fingerprint, so a change to weights/points/descriptions (not just the key set) triggers
        # it. save_rubric bumps the version, so this never overwrites an author-edited rubric.
        if int(rubric.get("version", 1) or 1) == 1 and rubric.get("seed_fp") != _SEED_FP:
            rubric = _seed()
            _write(rubric)
        return rubric


def save_rubric(dimensions: list[dict]) -> dict:
    """Persist an edited Kapture rubric → bumps version, returns the new rubric. Weights
    need not sum to 1 (the judge normalises at scoring time)."""
    cleaned = _clean_dimensions(dimensions)
    if not cleaned:
        raise ValueError("rubric must have at least one dimension")
    with _lock:
        current = _load() or _seed()
        rubric = {
            "version": int(current.get("version", 0)) + 1,
            "dimensions": cleaned,
            "updated_at": _now(),
        }
        _write(rubric)
        return rubric


# ── IMPORT FROM A STRUCTURED SHEET (Excel / CSV of factor · weight · description) ──
# Deterministic — no LLM. For a priority-factors SHEET the columns already are the rubric,
# so parse them exactly (precise weights) instead of asking a model to interpret them.
_LABEL_COLS = ("factor", "factors", "dimension", "dimensions", "label", "name", "criteria",
               "criterion", "parameter", "attribute", "quality parameter")
_WEIGHT_COLS = ("weight", "weightage", "weight %", "weight(%)", "priority", "importance",
                "score", "points", "%", "percentage", "value")
_DESC_COLS = ("description", "definition", "details", "detail", "notes", "note", "meaning",
              "what it measures", "guidance", "how to score")


def _rows_from_xlsx(raw: bytes) -> list[list[str]]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb.active
    rows = [["" if c is None else str(c).strip() for c in r] for r in ws.iter_rows(values_only=True)]
    wb.close()
    return rows


def _rows_from_csv(raw: bytes, delim: str = ",") -> list[list[str]]:
    text = raw.decode("utf-8-sig", errors="replace")
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), delimiter=delim)]


def dimensions_from_sheet(raw: bytes, filename: str) -> list[dict] | None:
    """Parse a factor/weight/description spreadsheet into rubric dimensions, deterministically.
    Returns dims (unnormalised weights — the judge sum-normalises) or None if the file isn't a
    recognisable factors sheet (caller then falls back to the LLM structure-from-text path)."""
    fn = (filename or "").lower()
    try:
        if fn.endswith((".xlsx", ".xlsm")):
            rows = _rows_from_xlsx(raw)
        elif fn.endswith(".tsv"):
            rows = _rows_from_csv(raw, "\t")
        elif fn.endswith((".csv", ".txt")):
            rows = _rows_from_csv(raw)
        else:
            return None
    except Exception:  # noqa: BLE001 — unreadable/binary → let the LLM path try
        return None

    # find the header row (one of the first ~10 rows containing a label-ish column)
    label_i = weight_i = desc_i = None
    header_idx = None
    for i, row in enumerate(rows[:10]):
        low = [c.lower() for c in row]
        li = next((j for j, c in enumerate(low) if c in _LABEL_COLS), None)
        if li is not None:
            header_idx = i
            label_i = li
            weight_i = next((j for j, c in enumerate(low) if c in _WEIGHT_COLS), None)
            desc_i = next((j for j, c in enumerate(low) if c in _DESC_COLS), None)
            break
    if header_idx is None:
        return None

    dims: list[dict] = []
    for row in rows[header_idx + 1:]:
        if label_i >= len(row):
            continue
        label = row[label_i].strip()
        if not label:
            continue
        weight = 1.0
        if weight_i is not None and weight_i < len(row):
            w = row[weight_i].replace("%", "").replace(",", "").strip()
            try:
                weight = float(w)
            except (TypeError, ValueError):
                weight = 1.0
        desc = row[desc_i].strip() if (desc_i is not None and desc_i < len(row)) else ""
        dims.append({"label": label, "weight": weight, "description": desc})
    return dims or None


def draft_from_dimensions(dims: list[dict]) -> dict:
    """Wrap sheet-parsed dimensions into a DRAFT rubric for review (never auto-saved)."""
    cleaned = _clean_dimensions(dims)
    if not cleaned:
        cleaned = [dict(d) for d in _SEED_DIMENSIONS]
    current = _load() or _seed()
    return {
        "version": int(current.get("version", 1) or 1),
        "dimensions": cleaned, "status": "draft", "updated_at": _now(),
        "note": "Imported from your priority-factors sheet — review the weights, then Save.",
    }


# ── STRUCTURE-FROM-TEXT (upload a QA-guidelines doc → weighted factors) ────────
# Cloned from knowledge/governance.py::structure_framework_from_text, retargeted at the
# rubric's flat {key,label,weight,description} dimension shape.
_SYSTEM = """You are the Audit-Rubric Compiler for a partner-support QA team. A QA lead uploads
a document describing HOW support tickets should be audited — the factors/criteria they grade on
and how heavily each matters. You STRUCTURE that document into a strict rubric of weighted
dimensions. You do NOT invent criteria — infer only what the document supports. Be precise."""

_PROMPT = """Convert the following audit-guidelines document into a scoring rubric. Return ONLY
JSON with EXACTLY this shape (the array MUST be present and non-empty):
{{
  "dimensions": [
    {{ "key": "<snake_case>", "label": "<short label>", "weight": <number 0-1>, "description": "<what this factor grades and how a good vs bad resolution looks>" }}
  ],
  "note": "<one sentence on what this rubric captures>"
}}

Rules:
- Each factor/criterion the document describes becomes one dimension with a snake_case key.
- Set each weight from the emphasis the document gives it (higher = more important). Weights need
  not sum to 1 — they are normalised at scoring time. If the document gives no weights, distribute
  them by apparent priority.
- Write a concrete description for each dimension so a judge can score 0.0-1.0 consistently.
- Infer only what the document supports; do not invent factors.

AUDIT-GUIDELINES DOCUMENT:
---
{text}
---
"""


def structure_rubric_from_text(text: str) -> dict:
    """LLM-structure an uploaded QA-guidelines document into a DRAFT rubric (weighted
    dimensions) for human review — never auto-saved. Robust JSON parse with one stricter
    retry. Reuses for_node("sop_compile") + json_mode + _parse_json, exactly like the
    governance framework's structure-from-text."""
    provider, model = llm_registry.for_node("sop_compile")
    prompt = _PROMPT.format(text=(text or "").strip())

    def _try(extra: str = "") -> dict:
        res = provider.generate(prompt + extra, model=model, node="sop_compile",
                                system=_SYSTEM, json_mode=True)
        parsed = _parse_json(res.text)
        if isinstance(parsed, list) and parsed:   # some runs wrap the object in an array
            parsed = parsed[0]
        return parsed if isinstance(parsed, dict) else {}

    parsed = _try()
    if not parsed or not parsed.get("dimensions"):
        parsed = _try("\n\nReturn STRICT valid JSON only — no prose, no markdown fences.")

    dims = _clean_dimensions(parsed.get("dimensions") if isinstance(parsed, dict) else [])
    if not dims:   # never hand back an empty rubric — fall back to the seed for review
        dims = [dict(d) for d in _SEED_DIMENSIONS]
    current = _load() or _seed()
    return {
        "version": int(current.get("version", 1) or 1),
        "dimensions": dims,
        "status": "draft",
        "updated_at": _now(),
        "note": (parsed.get("note") if isinstance(parsed, dict) else "")
                or "Structured from an uploaded document — review the factors + weights before saving.",
    }
