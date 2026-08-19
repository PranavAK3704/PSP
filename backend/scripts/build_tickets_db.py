"""Build a lean, PII-scrubbed, READ-ONLY tickets.db from a Kapture ticket export.

This is OFFLINE tooling — it runs once on a laptop to convert the 64 MB xlsx dump into a small
SQLite file the app reads read-only. The app NEVER writes to tickets.db at runtime (see
app/substrate/tickets_db.py, which opens it mode=ro). Rebuild only when a new export lands.

What it deliberately DROPS (partner PII / bulk / low-value):
  • Customer Name, Phone (both columns)   — personal PII, never stored
  • Ticket Remark (the transcript)         — PII-heavy free text, and huge
  • "describe issue" / "issue related to"  — free text that leaks names/phones/AWBs
  • FE ID, Enbolt ID, Transaction ID       — indirect person/txn identifiers

What it KEEPS (categorical + metrics — safe, and what an analytics widget needs):
  ticket_no, created_ts, resolved_ts, art_hours, status, sub_status, source,
  folder_l1, folder_l2, queue, current_queue, sub_type, sub_sub_type, hub_code,
  location_code, sla, payment_cycle

Usage:  python scripts/build_tickets_db.py ["<path to xlsx>"]
"""
from __future__ import annotations

import sqlite3
import sys
from datetime import datetime
from pathlib import Path

DEFAULT_XLSX = str(Path.home() / "Downloads" / "Ticket_Report_Raw DC Jul 31.xlsx")
OUT = Path(__file__).resolve().parents[1] / "data" / "tickets.db"

# header name  ->  column in tickets.db  (only these survive; everything else is dropped)
KEEP = {
    "Ticket No": "ticket_no",
    "Status": "status",
    "Sub Status": "sub_status",
    "Source": "source",
    "Disposition Folder Level 1": "folder_l1",
    "Disposition Folder Level 2": "folder_l2",
    "Queue": "queue",
    "Current Queue Name": "current_queue",
    "Sub Type": "sub_type",
    "Sub Sub Type": "sub_sub_type",
    "Hub Code": "hub_code",
    "Location Code": "location_code",
    "SLA": "sla",
    "Payment Cycle": "payment_cycle",
}
COLS = ["ticket_no", "created_ts", "resolved_ts", "art_hours", "status", "sub_status",
        "source", "folder_l1", "folder_l2", "queue", "current_queue", "sub_type",
        "sub_sub_type", "hub_code", "location_code", "sla", "payment_cycle"]


def _clean(v):
    """Whitespace-only cells are empty in this export — normalise them to NULL."""
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _ts(v):
    """'01-05-2026 00:02:32' -> ISO 'YYYY-MM-DD HH:MM:SS' (sortable). Best-effort."""
    s = _clean(v)
    if not s:
        return None
    for fmt in ("%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y %H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d %H:%M:%S")
        except ValueError:
            continue
    return s  # keep raw if unparseable — never crash the build


def main(xlsx: str) -> None:
    import openpyxl  # local import so the app never needs it at runtime

    print(f"reading {xlsx}")
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    header = list(next(it))
    idx = {}
    for i, h in enumerate(header):
        if h not in idx:            # first occurrence wins (dup column names in this export)
            idx[h] = i
    ci = {src: idx[src] for src in KEEP if src in idx}
    i_created = idx.get("Created Date")
    i_resolved = idx.get("Resolved Date")

    if OUT.exists():
        OUT.unlink()
    con = sqlite3.connect(OUT)
    con.execute(f"CREATE TABLE tickets ({', '.join(c + ' TEXT' for c in COLS)})")
    ins = f"INSERT INTO tickets ({', '.join(COLS)}) VALUES ({', '.join('?' for _ in COLS)})"

    batch, total = [], 0
    for row in it:
        vals = {dst: _clean(row[ci[src]]) for src, dst in KEEP.items() if src in ci}
        created = _ts(row[i_created]) if i_created is not None else None
        resolved = _ts(row[i_resolved]) if i_resolved is not None else None
        art = None
        if created and resolved:
            try:
                dt = (datetime.strptime(resolved, "%Y-%m-%d %H:%M:%S")
                      - datetime.strptime(created, "%Y-%m-%d %H:%M:%S"))
                art = round(dt.total_seconds() / 3600.0, 2)
            except ValueError:
                art = None
        rec = {**vals, "created_ts": created, "resolved_ts": resolved,
               "art_hours": art}
        batch.append(tuple(rec.get(c) for c in COLS))
        if len(batch) >= 5000:
            con.executemany(ins, batch); total += len(batch); batch = []
    if batch:
        con.executemany(ins, batch); total += len(batch)

    for col in ("hub_code", "status", "source", "folder_l1", "queue", "sla", "created_ts"):
        con.execute(f"CREATE INDEX ix_tickets_{col} ON tickets({col})")
    con.commit()
    n = con.execute("SELECT COUNT(*) FROM tickets").fetchone()[0]
    con.close(); wb.close()
    print(f"wrote {OUT}  ({total} rows inserted, {n} in table, {OUT.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX)
